"""PPTX import (LibreOffice to PDF, python-pptx for notes) and export."""

from __future__ import annotations

import asyncio
import io
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

PPTX_MAGIC = b"PK\x03\x04"  # PPTX is a zip container
MAX_PPTX_BYTES = 60 * 1024 * 1024
CONVERT_TIMEOUT_S = 180


def soffice_path() -> Optional[str]:
    return shutil.which("soffice") or shutil.which("libreoffice")


def extract_pptx_text(data: bytes) -> List[Dict[str, Any]]:
    """Per slide: title, visible text and speaker notes."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for slide in prs.slides:
        title = ""
        texts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if not text:
                continue
            if shape == getattr(slide.shapes, "title", None) or (not title and shape.is_placeholder and "title" in shape.name.lower()):
                title = text.split("\n")[0]
            else:
                texts.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append({"title": title, "text": "\n".join(texts), "notes": notes})
    return out


def _convert_sync(data: bytes) -> bytes:
    exe = soffice_path()
    if not exe:
        raise RuntimeError("PPTX import needs LibreOffice on the server (use the 'slides' Docker image).")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "deck.pptx"
        src.write_bytes(data)
        # A private profile dir lets several conversions run at once without LibreOffice's lock.
        profile = Path(tmp) / "lo-profile"
        subprocess.run(
            [exe, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
             "--convert-to", "pdf", "--outdir", tmp, str(src)],
            check=True, timeout=CONVERT_TIMEOUT_S, capture_output=True,
        )
        pdf = Path(tmp) / "deck.pdf"
        if not pdf.exists():
            raise RuntimeError("LibreOffice did not produce a PDF.")
        return pdf.read_bytes()


async def pptx_to_pdf(data: bytes) -> bytes:
    return await asyncio.to_thread(_convert_sync, data)


def export_pptx(deck: Dict[str, Any], images: Dict[int, Path]) -> bytes:
    """16:9 PPTX: picture slides for uploaded decks, title + bullets otherwise; notes kept."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank, content = prs.slide_layouts[6], prs.slide_layouts[1]
    for s in deck["slides"]:
        img = images.get(s["idx"])
        if img and img.exists():
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
        else:
            slide = prs.slides.add_slide(content)
            slide.shapes.title.text = s["title"] or " "
            body = slide.placeholders[1].text_frame
            body.clear()
            for i, b in enumerate(s["bullets"] or []):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = b
                p.font.size = Pt(28)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_sample_pptx(slides: Sequence[Dict[str, str]]) -> bytes:
    """Small PPTX with titles, bullets and notes (used by tests and the import self-check)."""
    return export_pptx({"slides": [{"idx": i, **s} for i, s in enumerate(slides)]}, {})
