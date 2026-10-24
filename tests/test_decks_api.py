import io
import json

import pytest

from backend.decks import pptx_tools
from backend.llm import router as router_mod

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
JPEG = b"\xff\xd8\xff\xe0" + b"\x00" * 64

OUTLINE = {
    "title": "FreightPilot Seed",
    "source": "outline",
    "slides": [
        {"title": "Shipping is broken", "bullets": ["Retailers overpay 30%"], "notes": "Open with the pain.", "time_budget_sec": 30},
        {"title": "FreightPilot", "bullets": ["Pool volume", "Auto-book carriers"], "notes": "", "time_budget_sec": 40},
        {"title": "The ask", "bullets": ["$1.5M seed"], "notes": "Close strong.", "time_budget_sec": 20},
    ],
}


def test_outline_deck_lifecycle(owner_client):
    c = owner_client
    deck = c.post("/api/decks", json=OUTLINE).json()
    assert [s["title"] for s in deck["slides"]] == ["Shipping is broken", "FreightPilot", "The ask"]
    assert c.get("/api/decks").json()["decks"][0]["slide_count"] == 3

    # Images: content-type and magic bytes must agree
    url = f"/api/decks/{deck['id']}/slides/1/image"
    assert c.put(url, content=PNG, headers={"content-type": "image/jpeg"}).status_code == 415
    assert c.put(url, content=b"<svg/>", headers={"content-type": "image/svg+xml"}).status_code == 415
    assert c.put(url, content=PNG, headers={"content-type": "image/png"}).status_code == 200
    assert c.get(url).content == PNG
    assert c.put(f"/api/decks/{deck['id']}/slides/9/image", content=JPEG, headers={"content-type": "image/jpeg"}).status_code == 404

    # Editing slide text keeps images attached by index and clears the stale critique
    edited = [dict(s, title=s["title"] + "!") for s in OUTLINE["slides"]]
    deck2 = c.put(f"/api/decks/{deck['id']}/slides", json={"slides": edited}).json()
    assert deck2["slides"][1]["has_image"] and deck2["slides"][0]["title"].endswith("!")

    assert c.patch(f"/api/decks/{deck['id']}", json={"title": "Renamed"}).json()["title"] == "Renamed"
    assert c.delete(f"/api/decks/{deck['id']}").status_code == 200
    assert c.get(f"/api/decks/{deck['id']}").status_code == 404


def test_decks_are_profile_scoped(owner_client):
    c = owner_client
    deck = c.post("/api/decks", json=OUTLINE).json()
    other = c.post("/api/profiles", json={"name": "Other"}).json()
    c.post(f"/api/profiles/{other['id']}/select", json={})
    assert c.get("/api/decks").json()["decks"] == []
    assert c.get(f"/api/decks/{deck['id']}").status_code == 404
    assert c.delete(f"/api/decks/{deck['id']}").status_code == 404


def test_export_pptx_keeps_titles_and_notes(owner_client):
    from pptx import Presentation

    c = owner_client
    deck = c.post("/api/decks", json=OUTLINE).json()
    res = c.get(f"/api/decks/{deck['id']}/export.pptx")
    assert res.status_code == 200 and "attachment" in res.headers["content-disposition"]
    prs = Presentation(io.BytesIO(res.content))
    assert len(prs.slides) == 3
    assert prs.slides[0].shapes.title.text == "Shipping is broken"
    assert prs.slides[2].notes_slide.notes_text_frame.text == "Close strong."


def test_pptx_text_extraction():
    data = pptx_tools.make_sample_pptx([{"title": "Hook", "bullets": ["Big pain"], "notes": "Say it slowly."}])
    slides = pptx_tools.extract_pptx_text(data)
    assert slides == [{"title": "Hook", "text": "Big pain", "notes": "Say it slowly."}]


@pytest.mark.skipif(not pptx_tools.soffice_path(), reason="LibreOffice not installed")
def test_convert_pptx_with_libreoffice(owner_client):
    c = owner_client
    data = pptx_tools.make_sample_pptx([
        {"title": "Problem", "bullets": ["Shipping costs"], "notes": "Pain first."},
        {"title": "Solution", "bullets": ["Pooling"], "notes": ""},
    ])
    res = c.post("/api/decks/convert-pptx", files={"file": ("deck.pptx", data, "application/octet-stream")})
    assert res.status_code == 200, res.text
    body = res.json()
    assert [s["title"] for s in body["slides"]] == ["Problem", "Solution"]
    assert body["slides"][0]["notes"] == "Pain first."
    pdf = c.get(body["pdf_url"])
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    # Not a PPTX
    assert c.post("/api/decks/convert-pptx", files={"file": ("x.pptx", b"hello", "application/octet-stream")}).status_code == 415


class FakeJSON:
    def __init__(self, payload):
        self.payload = payload
        self.calls = []

    async def generate_json(self, messages, model, **kw):
        self.calls.append(messages[-1].content)
        return json.dumps(self.payload)


def test_generate_deck_scales_budgets_to_requested_length(owner_client, monkeypatch):
    fake = FakeJSON({"title": "FreightPilot", "slides": [
        {"kind": "hook", "title": "Shipping eats margins", "bullets": ["30% overpay"], "notes": "n " * 50, "time_budget_sec": 10},
        {"kind": "problem", "title": "No leverage", "bullets": [], "notes": "n " * 50, "time_budget_sec": 10},
        {"kind": "solution", "title": "Pool and book", "bullets": [], "notes": "n " * 50, "time_budget_sec": 10},
        {"kind": "ask", "title": "$1.5M seed", "bullets": [], "notes": "n " * 50, "time_budget_sec": 10},
    ]})

    async def resolve(self, ref):
        return fake, "m"
    monkeypatch.setattr(router_mod.LLMRouter, "resolve_provider_and_model", resolve)

    res = owner_client.post("/api/decks/generate", json={"company": "FreightPilot", "one_liner": "Freight pooling for retailers", "minutes": 2})
    assert res.status_code == 200, res.text
    deck = res.json()
    assert deck["source"] == "ai" and deck["slides"][0]["kind"] == "hook"
    assert abs(sum(s["time_budget_sec"] for s in deck["slides"]) - 120) <= 4
    assert "120 seconds" in fake.calls[0]
    assert deck["brief"]["company"] == "FreightPilot"


def test_critique_keeps_only_real_slides(owner_client, monkeypatch):
    fake = FakeJSON({"overall": "Clear story.", "missing_sections": ["traction"], "slides": [
        {"slide_number": 2, "density": "too_dense", "issues": ["Two ideas"], "suggestion": "Split it"},
        {"slide_number": 9, "density": "ok", "issues": [], "suggestion": ""},
    ]})

    async def resolve(self, ref):
        return fake, "m"
    monkeypatch.setattr(router_mod.LLMRouter, "resolve_provider_and_model", resolve)
    deck = owner_client.post("/api/decks", json=OUTLINE).json()
    crit = owner_client.post(f"/api/decks/{deck['id']}/critique").json()
    assert [s["slide_number"] for s in crit["slides"]] == [2]
    assert len(crit["density"]) == 3 and crit["missing_sections"] == ["traction"]
    assert owner_client.get(f"/api/decks/{deck['id']}").json()["critique"]["overall"] == "Clear story."
